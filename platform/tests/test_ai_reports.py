import asyncio
import io
import zipfile

import pytest
from docx import Document
from pypdf import PdfReader
from talent_platform.reports import render_report


def test_docx_export_contains_query_table_and_embedded_chart_without_external_links():
    result = {'answer': '测试数据的部门统计', 'sql': 'SELECT department, total FROM safe_view',
              'columns': ['部门', '人数'], 'rows': [['研发', 12], ['采购', 8]],
              'chart': {'type': 'bar', 'url': 'file:///secret', 'data': [['伪造', 99999]]}}
    data = render_report(result, 'docx')
    document = Document(io.BytesIO(data))
    assert [[cell.text for cell in row.cells] for row in document.tables[0].rows] == [['部门', '人数'], ['研发', '12'], ['采购', '8']]
    assert len(document.inline_shapes) == 1
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        assert b'TargetMode="External"' not in b''.join(package.read(n) for n in package.namelist() if n.endswith('.rels'))
        assert b'99999' not in package.read('word/document.xml')


def test_pdf_export_contains_table_image_and_splits_very_long_cells():
    result = {'answer': '<img src="file:///secret"> & 测试', 'columns': ['部门', '人数'],
              'rows': [['研发' * 1300, 12], ['采购', -8]], 'sql': 'SELECT ...'}
    data = render_report(result, 'pdf')
    reader = PdfReader(io.BytesIO(data))
    assert len(reader.pages) >= 2
    assert any(page.images for page in reader.pages)
    assert sum('采购' in (page.extract_text() or '') for page in reader.pages) >= 1
    assert sum((page.extract_text() or '').count('部门') for page in reader.pages) <= len(reader.pages)
    assert all(len(page.extract_text() or '') > 80 for page in reader.pages)


def test_non_numeric_rows_cannot_produce_a_misleading_chart():
    result = {'answer': '类别明细', 'columns': ['姓名', '状态'], 'rows': [['张三', '在岗']]}
    doc = Document(io.BytesIO(render_report(result, 'docx')))
    assert len(doc.tables) == 1
    assert len(doc.inline_shapes) == 0


def test_analysis_clarification_returns_session_without_running_sql():
    from talent_platform.analysis import AnalyticsEngine
    seen = []
    async def model(request, actor):
        seen.append(request)
        return {'clarification': '请确认统计月份和指标'}
    async def reader(sql, dialect):
        pytest.fail('clarification must not execute SQL')
    result = asyncio.run(AnalyticsEngine(model=model, runner=reader).execute({
        'question': '它怎么样', 'datasource': 'postgres', 'allowed_org_ids': ['o'], 'allowed_fields': ['name'],
        'session_id': 's', 'history': [{'role': 'user', 'content': '上个月的人才'}]},
        actor={'id': 'u', 'org_id': 'o', 'role': 'hr'}))
    assert result == {'needs_clarification': True, 'question': '请确认统计月份和指标', 'session_id': 's'}
    assert seen[0]['history'] == [{'role': 'user', 'content': '上个月的人才'}]


def test_analysis_uses_line_chart_for_ordered_time_series_and_numeric_decimal_strings():
    from talent_platform.analysis import AnalyticsEngine
    result = AnalyticsEngine._chart(['月份', '人数'], [['2026-01', '12.50'], ['2026-02', '13.75']])
    assert result['type'] == 'line'
    assert result['data'] == [['2026-01', 12.5], ['2026-02', 13.75]]
    assert AnalyticsEngine._chart(['类别', '值'], [['A', True]]) is None
