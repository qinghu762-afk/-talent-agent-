import hashlib
import json
import runpy
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader


def generator():
    return runpy.run_path(str(Path(__file__).resolve().parents[1] / 'tools' / 'generate_acceptance_fixtures.py'))['generate']


def test_fixture_generator_creates_safe_labelled_multiformat_corpus(tmp_path):
    output = tmp_path / 'fixtures'
    manifest = generator()(output)
    records = manifest['fixtures']
    assert len(records) >= 30
    assert manifest['synthetic'] is True
    assert {record['expected_quality'] for record in records} == {'complete', 'partial', 'fail'}
    assert {'txt', 'docx', 'pdf_native', 'png_scan', 'jpg_scan', 'pdf_scan', 'xlsx', 'pdf_corrupt'} <= {record['type'] for record in records}
    for record in records:
        assert Path(record['filename']).name == record['filename']
        assert record['synthetic'] is True
        path = output / record['filename']
        assert path.is_file()
        assert hashlib.sha256(path.read_bytes()).hexdigest() == record['sha256']
        if record['expected_quality'] == 'complete':
            assert record['expected_profile']['email'].endswith('@example.test')
        assert record['expected_phrases'] or record['expected_quality'] == 'fail'
    assert json.loads((output / 'fixture-manifest.json').read_text(encoding='utf-8')) == manifest
    native = next(record for record in records if record['type'] == 'pdf_native' and record['layout'] == 'multipage')
    assert len(PdfReader(output / native['filename']).pages) >= 2
    image_pdf = next(record for record in records if record['type'] == 'pdf_scan')
    pages = PdfReader(output / image_pdf['filename']).pages
    assert any(page.images for page in pages)
    assert not ''.join(page.extract_text() or '' for page in pages).strip()
    two_column = next(record for record in records if record['type'] == 'docx' and record['layout'] == 'two_column')
    document = Document(output / two_column['filename'])
    assert 'w:num="2"' in document.sections[0]._sectPr.xml
    tabular = next(record for record in records if record['type'] == 'docx' and record['layout'] == 'table')
    assert Document(output / tabular['filename']).tables
    workbook = next(record for record in records if record['type'] == 'xlsx')
    assert load_workbook(output / workbook['filename'], data_only=True).active['A1'].value.startswith('合成')
    for record in records:
        if record['type'] in {'png_scan', 'jpg_scan'}:
            with Image.open(output / record['filename']) as image:
                assert image.width >= 1200 and image.height >= 1500


def test_fixture_generator_refuses_existing_material_without_overwriting(tmp_path):
    output = tmp_path / 'existing'
    output.mkdir()
    original = output / 'keep.txt'
    original.write_text('existing material', encoding='utf-8')
    with pytest.raises(FileExistsError):
        generator()(output)
    assert original.read_text(encoding='utf-8') == 'existing material'
