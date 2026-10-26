"""Durable PostgreSQL outbox worker. Redis never owns business facts."""
import asyncio
import hashlib
import json
import logging
from contextlib import suppress
from datetime import timedelta
from pathlib import Path

import httpx
from sqlalchemy import and_, or_, select
from sqlalchemy.exc import OperationalError

from .config import settings
from .db import Document, Record, Resource, Run, Stores, User, now
from .persistence import persist_resume
from .security import redact, resolved_user, workflow_permitted, authorization_snapshot
from .tasks import enqueue, event

log = logging.getLogger("talent.worker")
TERMINAL = {"succeeded", "failed", "cancelled", "waiting_approval"}


class Worker:
    def __init__(self, config=None, stores=None, executor=None):
        self.config = config or settings()
        self.stores = stores or Stores(self.config)
        self.executor = executor or self.call_ai

    async def call_ai(self, kind, payload, run_id, actor):
        url = self.config.analytics_service_url if kind == "analysis" else self.config.ai_service_url
        async with httpx.AsyncClient(timeout=httpx.Timeout(900, connect=10)) as client:
            r = await client.post(url + "/execute", json={"kind": kind, "payload": payload, "run_id": run_id, "actor": actor}, headers={"X-Service-Key": self.config.internal_service_key})
            r.raise_for_status()
            value = r.json()
            return value.get("data", value)

    @staticmethod
    def memory_scope(who):
        return f"{who.get('role')}:{who.get('auth_version')}:{who.get('policy_version')}"

    async def memory_cache(self, history, who, session_id, channel):
        if not self.stores.redis:
            return history
        from redis.exceptions import RedisError
        encoded = json.dumps(history, ensure_ascii=False)
        digest = hashlib.sha256(encoded.encode()).hexdigest()
        key = f"memory:{who['org_id']}:{who['id']}:{self.memory_scope(who)}:{channel}:{session_id}:{digest}"
        try:
            cached = await self.stores.redis.get(key)
            if cached:
                return json.loads(cached)
            await self.stores.redis.setex(key, 900, encoded)
        except RedisError:
            pass
        return history

    async def claim(self, ident=None):
        async with self.stores.facts() as db:
            q = select(Run).where(or_(and_(Run.status == "queued", Run.available_at <= now()), and_(Run.status == "running", Run.lease_until < now()))).order_by(Run.priority, Run.created_at).with_for_update(skip_locked=True).limit(1)
            if ident:
                q = q.where(Run.id == ident)
            run = await db.scalar(q)
            if not run:
                return None
            run.status = "running"
            run.attempt += 1
            run.lease_until = now() + timedelta(seconds=self.config.worker_lease_seconds)
            run.updated_at = now()
            await event(db, run)
            await db.commit()
            return run

    async def heartbeat(self, ident, attempt):
        while True:
            await asyncio.sleep(20)
            async with self.stores.facts() as db:
                run = await db.get(Run, ident)
                if not run or run.status != "running" or run.attempt != attempt:
                    return
                run.lease_until = now() + timedelta(seconds=self.config.worker_lease_seconds)
                await db.commit()

    async def current_actor(self, run):
        async with self.stores.business() as db:
            user = await db.get(User, run.owner_id)
            if not user or not user.active or user.org_id != run.actor.get("org_id") or user.role != run.actor.get("role"):
                raise PermissionError("任务账号或权限已变更")
            who = await resolved_user(db, user)
            if any(run.actor.get(k) is not None and run.actor.get(k) != who.get(k) for k in ("auth_version", "policy_version")):
                raise PermissionError("任务授权已变更")
            return who

    async def record(self, who, kind, ident):
        from .api import resource
        return await resource(self.stores, who, kind, ident)

    async def records(self, who, kind):
        from .api import listed
        return await listed(self.stores, who, kind)

    async def snapshot(self, run, who):
        p = run.payload
        kind = run.kind
        if not workflow_permitted(who, kind):
            raise PermissionError('当前角色不能执行该工作流')
        if who.get("permissions") is not None:
            from fnmatch import fnmatchcase
            operation = {"resume": "POST:talents", "batch": "POST:talents", "job": "POST:jobs", "job_import": "POST:jobs", "match": "POST:jobs", "search": "POST:talents", "chat": "POST:chat", "analysis": "POST:analysis", "assessment": "POST:assessments", "training": "POST:training-plans"}.get(kind)
            if operation and not any(fnmatchcase(operation, perm) for perm in who["permissions"]):
                raise PermissionError("自定义角色未授权此工作流")
        if kind == "resume":
            async with self.stores.facts() as db:
                existing = await db.scalar(select(Document).where(Document.org_id == who["org_id"], Document.sha256 == p["sha256"]))
                if existing and existing.talent_id:
                    return {"_cached": {"talent_id": existing.talent_id, "document_id": existing.id, "reused": True, 'report': existing.model_versions.get('resume_workflow', {})}}
            return {"text": p.get("text") or await self.parse(p["path"], p["filename"])}
        if kind == "job":
            return {"description": p["description"]}
        if kind in {"index", "duplicates"}:
            t = await self.record(who, "talents", p["talent_id"])
            result = {"talent": self.talent_snapshot(t)}
            if kind == "duplicates":
                result["allowed_ids"] = [row.id for row in await self.records(who, "talents") if row.org_id == t.org_id and row.indexed_revision == row.revision and row.id != t.id]
            return result
        if kind == "match":
            job = await self.record(who, "jobs", p["job_id"])
            talents = [t for t in await self.records(who, "talents") if t.org_id == job.org_id and t.revision == t.indexed_revision and not t.data.get("deleted")]
            return {"job": {"id": job.id, "org_id": job.org_id, "title": job.data.get("title", ""), "requirements": job.data.get("requirements", {}), "revision": job.revision}, "talents": [self.talent_snapshot(t) | {"profile": redact(t.data.get("profile", {}), who)} for t in talents], "allowed_ids": [t.id for t in talents]}
        if kind == "search":
            talents = [t for t in await self.records(who, "talents") if t.revision == t.indexed_revision]
            return {"query": p["query"], "talents": [self.talent_snapshot(t) | {"profile": redact(t.data.get("profile", {}), who)} for t in talents], "allowed_ids": [t.id for t in talents]}
        if kind == "knowledge":
            doc = await self.record(who, "knowledge", p["document_id"])
            return {"document": {**doc.data, "id": doc.id, "org_id": doc.org_id, "owner_id": doc.owner_id, "revision": doc.revision}}
        if kind == "chat":
            docs = [d for d in await self.records(who, "knowledge") if d.indexed_revision == d.revision and not d.data.get("deleted")]
            documents = [{**d.data, "id": d.id, "org_id": d.org_id, "owner_id": d.owner_id, "revision": d.revision, "indexed_revision": d.indexed_revision, "index_model": d.data.get("index_model"), "kind": "knowledge"} for d in docs]
            if p.get("talent_id") or who["role"] in {"admin", "hr", "manager"}:
                talents = await self.records(who, "talents")
                for t in talents:
                    if t.indexed_revision != t.revision or p.get("talent_id") and t.id != p["talent_id"]:
                        continue
                    documents.append({"id": t.id, "org_id": t.org_id, "owner_id": t.owner_id, "revision": t.revision, "indexed_revision": t.indexed_revision, "index_model": t.data.get("index_model"), "kind": "talent", "scope": "org", "title": (t.data.get("profile", {}).get("name") or "人才") + "档案", "text": json.dumps(redact(t.data.get("profile", {}), who), ensure_ascii=False)})
            current_docs = {d['id']: d for d in documents}
            memories = [m for m in await self.records(who, "memories") if m.owner_id == who["id"] and m.data.get("channel", "chat") == "chat" and m.data.get("session_id") == p["session_id"] and m.data.get("scope_version") == self.memory_scope(who)
                        and all(s.get('source_id') in current_docs and s.get('revision') == current_docs[s['source_id']].get('revision') for s in m.data.get('source_snapshots', []))]
            history = [x for m in sorted(memories, key=lambda m: m.created_at) for x in m.data.get("messages", [])][-12:]
            history = await self.memory_cache(history, who, p["session_id"], "chat")
            return {**p, "documents": documents, "history": history}
        if kind == "assessment":
            exam = await self.record(who, "assessments", p["assessment_id"])
            # Employee cannot use question API, but worker reads only exam-authorized questions.
            qs = exam.data.get("question_snapshot") or [{**q.data, "id": q.id, "revision": q.revision} for q in [await self.record({**who, "role": "hr"}, "questions", q) for q in exam.data["question_ids"]]]
            talent = await self.record(who, "talents", exam.data["talent_id"])
            answers = p.get("answers")
            if answers is None:
                async with self.stores.facts() as db:
                    submitted = await db.get(Run, exam.data.get("run_id", ""))
                    if not submitted or submitted.kind != "assessment" or not submitted.payload.get("answers"):
                        raise ValueError("测评尚无已提交的可复核答案")
                    answers = submitted.payload["answers"]
            return {"questions": qs, "answers": answers, "events": p.get("events", []), "talent": self.talent_snapshot(talent)}
        if kind == "training":
            talent = await self.record(who, "talents", p["talent_id"])
            assessment = {}
            if p.get("assessment_id"):
                exam = await self.record(who, "assessments", p["assessment_id"])
                async with self.stores.facts() as db:
                    scored = await db.get(Run, exam.data.get("run_id", ""))
                    if scored and scored.status == "succeeded":
                        assessment = scored.result
            courses = await self.records(who, "courses")
            return {"talent": self.talent_snapshot(talent), "assessment": assessment, "courses": [{**c.data, "id": c.id} for c in courses if not c.data.get("deleted")]}
        if kind == "analysis":
            if who["role"] not in {"admin", "manager"}:
                raise PermissionError("无数据决策权限")
            if who.get("data_scope") == "self":
                raise PermissionError("个人数据范围不能执行组织统计")
            orgs = [who["org_id"]]
            if who["role"] == "admin":
                orgs = list({r.org_id for k in ("talents", "assessments", "training-plans") for r in await self.records(who, k)} | {r.id for r in await self.records(who, "organizations")} | {who["org_id"]})
            fields = {"analytics_talents": ["id", "org_id", "name", "location", "experience_months", "completeness", "revision"], "analytics_jobs": ["id", "org_id", "title", "revision"], "analytics_assessments": ["id", "org_id", "talent_id", "title", "status", "total", "revision"], "analytics_learning": ["id", "org_id", "talent_id", "title", "status", "progress", "revision"]}
            tables = ["analytics_jobs"] if p.get("datasource") == "mysql" else ["analytics_talents", "analytics_assessments", "analytics_learning"]
            memories = [m for m in await self.records(who, "memories") if m.owner_id == who["id"] and m.data.get("channel") == "analysis" and m.data.get("scope_version") == self.memory_scope(who) and m.data.get("session_id") == p["session_id"] and m.data.get("datasource") == p.get("datasource")]
            history = [item for memory in sorted(memories, key=lambda m: m.created_at) for item in memory.data.get("messages", [])][-12:]
            history = await self.memory_cache(history, who, p["session_id"], "analysis")
            return {**p, "history": history, "allowed_org_ids": orgs, "allowed_fields": [table + "." + field for table in tables for field in fields[table] if field not in who.get("denied_fields", [])], "allowed_tables": tables}
        if kind == "delete_index":
            return p
        raise ValueError("未知工作流类型")

    @staticmethod
    def talent_snapshot(t):
        return {"id": t.id, "org_id": t.org_id, "revision": t.revision, "indexed_revision": t.indexed_revision, "index_model": t.data.get("index_model"), "profile": t.data.get("profile", {})}

    async def parse(self, path, filename):
        content = Path(path).read_bytes()
        suffix = Path(filename).suffix.lower()
        if suffix == ".txt":
            for encoding in ("utf-8-sig", "gb18030"):
                try:
                    return content.decode(encoding)
                except UnicodeDecodeError:
                    continue
            raise ValueError("文本编码无法解析")
        if suffix == ".docx":
            from docx import Document as WordDocument
            doc = WordDocument(path)
            return "\n".join([p.text for p in doc.paragraphs] + [" | ".join(c.text for c in r.cells) for t in doc.tables for r in t.rows])
        if suffix in {".xlsx", ".xls"}:
            import pandas as pd
            sheets = await asyncio.to_thread(pd.read_excel, path, sheet_name=None, header=None)
            return "\n\n".join(str(name) + "\n" + '\n'.join(
                ': '.join(str(cell) for cell in row if str(cell).strip())
                for row in frame.fillna('').itertuples(index=False, name=None)) for name, frame in sheets.items())
        if suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(path)
            pages = [page.extract_text() or '' for page in reader.pages]
            native = "\n".join(f"[第{i+1}页]\n{text}" for i, text in enumerate(pages))
            if pages and all(sum(c.isalnum() for c in text) >= 20 for text in pages):
                return native
        async with httpx.AsyncClient(timeout=600) as client:
            r = await client.post(self.config.parser_url + "/parse", files={"file": (filename, content)}, headers={"X-Service-Key": self.config.internal_service_key})
            r.raise_for_status()
            result = r.json()
            return result.get("data", result)["text"]

    async def apply(self, run, who, result, snapshot):
        from .api import put
        p = run.payload
        if run.kind in {"match", "search"}:
            key = "items" if run.kind == "search" else "candidates"
            valid = []
            for item in result.get(key, []):
                try:
                    latest = await self.record(who, "talents", item["talent_id"])
                    if latest.revision == item.get("revision") == latest.indexed_revision:
                        valid.append(item)
                except Exception as exc:
                    from fastapi import HTTPException
                    if not isinstance(exc, HTTPException):
                        raise
            result = {**result, key: valid}
            if run.kind == "search":
                result["total"] = len(valid)
        if run.kind == "resume":
            return await persist_resume(self.stores, who, p, result)
        if run.kind in {"job", "job_import"}:
            jobs = await self.records(who, "jobs")
            row = next((j for j in jobs if j.data.get("created_by_run") == run.id), None)
            if not row:
                row = await put(self.stores, who, "jobs", {"title": p.get("title", "新岗位"), "description": p["description"], "requirements": result["requirements"], "created_by_run": run.id, "versions": result.get("versions", {})})
            return {**result, "job_id": row.id}
        if run.kind in {"index", "knowledge"}:
            ident = p["talent_id"] if run.kind == "index" else p["document_id"]
            async with self.stores.facts() as db:
                row = await db.scalar(select(Record).where(Record.id == ident).with_for_update())
                version = snapshot["talent" if run.kind == "index" else "document"]["revision"]
                if row and row.revision == version and not row.data.get("deleted"):
                    row.indexed_revision = version
                    row.data = {**row.data, "index_model": result.get("index_model", self.config.embedding_model)}
                await db.commit()
            if run.kind == "index":
                await enqueue(self.stores, who, "duplicates", {"talent_id": ident}, 25)
            return result
        if run.kind == "duplicates":
            issues = await self.records(who, "issues")
            if result.get("suspected_duplicates") and not any(r.data.get("created_by_run") == run.id for r in issues):
                await put(self.stores, who, "issues", {"talent_id": p["talent_id"], "type": "suspected_duplicate", "suspected_duplicates": result["suspected_duplicates"], "created_by_run": run.id, "auto_merged": False})
            return result
        if run.kind == "assessment":
            exam = await self.record(who, "assessments", p["assessment_id"])
            existing = [r for r in await self.records(who, "assessment-results") if r.data.get("run_id") == run.id]
            if not existing:
                await put(self.stores, who, "assessment-results", {"run_id": run.id, "assessment_id": exam.id, "talent_id": exam.data["talent_id"], "answers": snapshot["answers"], "question_snapshot": snapshot["questions"], "result": result})
            await put(self.stores, who, "assessments", {"status": "waiting_rating", "total": result.get("total"), "result": result}, exam.id)
            approvals = await self.records({**who, "role": "admin"}, "approvals")
            approval = next((a for a in approvals if a.data.get("run_id") == run.id), None)
            if not approval:
                approval = await put(self.stores, who, "approvals", {"kind": "assessment_rating", "title": "测评评级复核", "target_id": exam.id, "run_id": run.id, "status": "pending", "inference": True, "risks": result.get("risks", [])})
            from .notifications import notify
            await notify(self.stores, key=approval.id, org_id=exam.org_id, title="测评结果待评级复核", message="AI 评分与风险记录已生成，请核对依据后审批。", target="approvals", managers=True)
            return {**result, "approval_id": approval.id}
        if run.kind == "training":
            plans = await self.records(who, "training-plans")
            plan = next((p for p in plans if p.data.get("run_id") == run.id), None)
            if not plan:
                plan = await put(self.stores, who, "training-plans", {**result, "talent_id": p["talent_id"], "status": "assigned", "due_at": (now() + timedelta(days=14)).isoformat(), "run_id": run.id, "baseline_assessment_id": p.get("assessment_id")})
            from .notifications import notify
            await notify(self.stores, key=run.id, org_id=plan.org_id, title="新的个性化学习计划", message=result.get("title", "请进入学习中心查看课程"), talent_id=p["talent_id"], target="training-plans")
            return {**result, "plan_id": plan.id}
        if run.kind == "analysis":
            memories = await self.records(who, "memories")
            if not any(m.data.get("run_id") == run.id for m in memories):
                await put(self.stores, who, "memories", {"run_id": run.id, "channel": "analysis", "datasource": p.get("datasource"), "scope_version": self.memory_scope(who), "session_id": p["session_id"], "scope": "private", "messages": [{"role": "user", "content": p["question"]}, {"role": "assistant", "content": result.get("question") or result.get("answer", "")}]})
            if result.get("needs_clarification"):
                return result
            reports = await self.records(who, "reports")
            report = next((r for r in reports if r.data.get("run_id") == run.id), None)
            if not report:
                report = await put(self.stores, who, "reports", {**result, "run_id": run.id, "question": p["question"], "title": "数据决策报告", "inference": True, 'scope': 'private', 'authorization': authorization_snapshot(who)})
            return {**result, "report_id": report.id}
        if run.kind == "chat":
            memories = await self.records(who, "memories")
            if not any(m.data.get("run_id") == run.id for m in memories):
                await put(self.stores, who, "memories", {"run_id": run.id, "session_id": p["session_id"], "scope_version": self.memory_scope(who), "source_snapshots": [{"source_id": c.get("source_id"), "revision": c.get("revision")} for c in result.get("citations", [])], "scope": "private", "messages": [{"role": "user", "content": p["question"]}, {"role": "assistant", "content": result.get("answer", "")} ]})
        return result

    async def batch(self, run, who):
        async with self.stores.facts() as db:
            # Resuming a batch discovers previously persisted children, never recreates them.
            previous = (await db.scalars(select(Run).where(Run.org_id == who["org_id"], Run.kind == "resume"))).all()
            children = {r.payload.get("batch_position"): r for r in previous if r.payload.get("parent_id") == run.id}
            for i, file in enumerate(run.payload["files"]):
                if i not in children:
                    child = Run(org_id=who["org_id"], owner_id=who["id"], actor=who, kind="resume", payload={**file, "parent_id": run.id, "batch_position": i}, priority=30)
                    db.add(child)
                    await db.flush()
                    await event(db, child)
                    children[i] = child
            await db.commit()
        settled = sum(child.status in TERMINAL for child in children.values())
        return {"files": [{"id": children[i].id, "filename": f["filename"], "status": children[i].status, "error": children[i].error} for i, f in enumerate(run.payload["files"])], "count": len(children), "settled": settled, "failed": sum(c.status == "failed" for c in children.values())}

    async def skill(self, run, who):
        from jsonschema import validate
        skill = await self.record(who, "skills", run.payload["skill_id"])
        if not skill.data.get("enabled", True):
            raise ValueError("技能已停用")
        parameters = run.payload.get("parameters", {})
        validate(parameters, skill.data.get("input_schema") or {"type": "object"})
        handler = skill.data["handler"]
        if not workflow_permitted(who, handler):
            raise PermissionError('当前角色不能执行该技能处理器')
        allowed = {"resume": {"text"}, "job": {"title", "description"}, "match": {"job_id"}, "assessment": {"assessment_id"}, "chat": {"question", "session_id"}, "training": {"talent_id", "assessment_id"}, "analysis": {"question", "datasource", "session_id"}}
        if handler not in allowed:
            raise ValueError("技能处理器未经注册")
        payload = {k: v for k, v in parameters.items() if k in allowed[handler]}
        if handler == "resume":
            from .uploads import save_file
            digest, path = save_file(self.config.data_dir, "skill-resume.txt", payload.get("text", "").encode())
            payload.update({"filename": "skill-resume.txt", "sha256": digest, "path": path})
        if handler in {"chat", "analysis"}:
            payload.setdefault("session_id", run.id)
        payload["__event_depth"] = int(run.payload.get("__event_depth", 0))
        # Child ID reuse makes skill execution idempotent across restarts.
        async with self.stores.facts() as db:
            child = await db.scalar(select(Run).where(Run.id == run.id + "-child"))
            if not child:
                import uuid
                ident = str(uuid.uuid5(uuid.UUID(run.id), "skill-child"))
                child = await db.get(Run, ident)
                if not child:
                    child = Run(id=ident, org_id=who["org_id"], owner_id=who["id"], actor=who, kind=handler, payload=payload, priority=5)
                    db.add(child)
                    await db.flush()
                    await event(db, child)
                    await db.commit()
        return {"child_run_id": child.id, "skill_version": skill.data.get("version", "1")}

    async def run_once(self, ident=None):
        run = await self.claim(ident)
        if not run:
            return False
        heartbeat = asyncio.create_task(self.heartbeat(run.id, run.attempt))
        try:
            who = await self.current_actor(run)
            if not workflow_permitted(who, run.kind):
                raise PermissionError('当前角色不能执行该工作流')
            if run.kind == "batch":
                result = await self.batch(run, who)
            elif run.kind == "skill":
                result = await self.skill(run, who)
            elif run.kind == "knowledge_import":
                from .api import put
                existing = [r for r in await self.records(who, "knowledge") if r.data.get("created_by_run") == run.id]
                if existing:
                    row = existing[0]
                else:
                    parsed = await self.parse(run.payload['path'], run.payload['filename'])
                    if not await self.active(run):
                        return True
                    who = await self.current_actor(run)
                    row = await put(self.stores, who, 'knowledge', {'title': run.payload['title'], 'scope': run.payload['scope'], 'text': parsed, 'created_by_run': run.id})
                result = {"document_id": row.id, "index_run_id": getattr(row, "_index_run_id", None)}
            elif run.kind == "job_import":
                description = await self.parse(run.payload["path"], run.payload["filename"])
                result = await self.executor("job", {"description": description}, run.id, who)
                if not await self.active(run):
                    return True
                who = await self.current_actor(run)
                run.payload = {**run.payload, "description": description}
                result = await self.apply(run, who, result, {"description": description})
            else:
                snapshot = run.payload.get("__snapshot") or await self.snapshot(run, who)
                if "_cached" in snapshot:
                    result = snapshot["_cached"]
                else:
                    async with self.stores.facts() as db:
                        current = await db.get(Run, run.id)
                        current.payload = {**current.payload, "__snapshot": snapshot}
                        await db.commit()
                    result = await self.executor(run.kind, snapshot, run.id, who)
                    async with self.stores.facts() as db:
                        current = await db.get(Run, run.id)
                        if current.status != "running" or current.attempt != run.attempt:
                            return True
                    who = await self.current_actor(run)
                    result = await self.apply(run, who, result, snapshot)
            async with self.stores.facts() as db:
                current = await db.scalar(select(Run).where(Run.id == run.id).with_for_update())
                if current.status != "running" or current.attempt != run.attempt:
                    return True
                pending_children = run.kind == "batch" and result.get("settled", 0) < result.get("count", 0)
                current.status = "running" if pending_children else "waiting_approval" if run.kind == "assessment" else "succeeded"
                if pending_children:
                    current.lease_until = now() + timedelta(seconds=2)
                current.result = result
                current.versions = {**(result.get("versions") or result.get('report', {}).get('versions', {})), "workflow": "talent-platform-v1", "scoring": "40-25-25-10-projects-v2",
                                    "job": snapshot.get("job") if run.kind in {"match", "search"} else None,
                                    "talents": [{k: t.get(k) for k in ("id", "revision", "indexed_revision", "index_model")} for t in snapshot.get("talents", [])] if run.kind in {"match", "search"} else []} if run.kind not in {"batch", "skill", "job_import", "knowledge_import"} else result.get("versions", {})
                current.progress = 100 * result["settled"] / max(1, result["count"]) if pending_children else 100
                current.error = None
                current.updated_at = now()
                await event(db, current)
                if current.status == "succeeded" and int(run.payload.get("__event_depth", 0)) < 4:
                    await self.emit_automations(db, run, result)
                await db.commit()
        except asyncio.CancelledError:
            # A graceful application restart is not a model failure. Retain the
            # frozen snapshot/checkpoint and release only our current lease.
            async with self.stores.facts() as db:
                current = await db.scalar(select(Run).where(Run.id == run.id).with_for_update())
                if current and current.status == 'running' and current.attempt == run.attempt:
                    current.status = 'queued'
                    current.attempt = max(0, current.attempt - 1)
                    current.available_at = now()
                    current.lease_until = now()
                    await event(db, current)
                    await db.commit()
            raise
        except Exception as exc:
            # Error text deliberately excludes request bodies and model responses.
            temporary = isinstance(exc, (httpx.TimeoutException, httpx.ConnectError, ConnectionError, TimeoutError, OperationalError)) or isinstance(exc, httpx.HTTPStatusError) and exc.response.status_code >= 500
            async with self.stores.facts() as db:
                current = await db.get(Run, run.id)
                if current.status == "running" and current.attempt == run.attempt:
                    current.status = "queued" if temporary and run.attempt <= 3 else "failed"
                    reason = "临时服务故障，等待重试" if current.status == "queued" else "服务重试预算已耗尽" if temporary else "输入文档或结构化结果未通过校验"
                    if isinstance(exc, PermissionError):
                        reason = "账号、组织或操作权限已变更，任务已停止"
                    if isinstance(exc, FileNotFoundError):
                        reason = "原文件不可用，请重新上传"
                    if isinstance(exc, httpx.HTTPStatusError) and exc.response.status_code == 422:
                        reason = "原文或模型输出没有可验证的有效结构，请查看原文件和字段依据"
                    current.error = f"{type(exc).__name__}：{reason}"
                    current.available_at = now() + timedelta(seconds=2 ** run.attempt)
                    await event(db, current)
                    await db.commit()
            log.warning("task %s %s failed (%s)", run.id, run.kind, type(exc).__name__)
        finally:
            heartbeat.cancel()
            with suppress(asyncio.CancelledError):
                await heartbeat
        return True

    async def active(self, run):
        async with self.stores.facts() as db:
            current = await db.get(Run, run.id)
            return bool(current and current.status == 'running' and current.attempt == run.attempt)

    async def emit_automations(self, facts, run, result):
        import uuid
        async with self.stores.business() as db:
            rows = (await db.scalars(select(Resource).where(Resource.kind == "automations", Resource.org_id == run.org_id))).all()
            for automation in rows:
                d = automation.data
                if not d.get("enabled", True) or d.get("trigger") != "event" or d.get("event_kind") != run.kind:
                    continue
                user = await db.get(User, automation.owner_id)
                if not user or not user.active or user.org_id != run.org_id:
                    continue
                ident = str(uuid.uuid5(uuid.UUID(automation.id), run.id))
                if await facts.get(Run, ident):
                    continue
                who = await resolved_user(db, user)
                values = {**run.payload, **result}
                def expand(v):
                    if isinstance(v, str) and v.startswith("$event."):
                        key = v.removeprefix("$event.")
                        return values.get(key) if key in {"talent_id", "assessment_id", "job_id", "plan_id", "report_id", "document_id"} else None
                    if isinstance(v, dict):
                        return {k: expand(x) for k, x in v.items()}
                    if isinstance(v, list):
                        return [expand(x) for x in v]
                    return v
                child = Run(id=ident, org_id=who["org_id"], owner_id=who["id"], actor=who, kind="skill", payload={"skill_id": d["skill_id"], "parameters": expand(d.get("parameters", {})), "__event_depth": int(run.payload.get("__event_depth", 0)) + 1}, priority=15)
                facts.add(child)
                await facts.flush()
                await event(facts, child)

    async def schedule(self):
        """Idempotent scheduled notifications and registered skill triggers."""
        async with self.stores.business() as db:
            automations = (await db.scalars(select(Resource).where(Resource.kind == "automations").with_for_update(skip_locked=True))).all()
            for automation in automations:
                d = automation.data
                if d.get("trigger") == "event" or not d.get("enabled", True) or d.get("next_at", "") > now().isoformat():
                    continue
                user = await db.get(User, automation.owner_id)
                if not user or not user.active:
                    continue
                who = await resolved_user(db, user)
                slot = int(now().timestamp()) // max(60, int(d.get("interval_seconds", 3600)))
                import uuid
                run_id = str(uuid.uuid5(uuid.UUID(automation.id), str(slot)))
                async with self.stores.facts() as facts:
                    if not await facts.get(Run, run_id):
                        run = Run(id=run_id, org_id=who["org_id"], owner_id=who["id"], actor=who, kind="skill", payload={"skill_id": d["skill_id"], "parameters": d.get("parameters", {})}, priority=15)
                        facts.add(run)
                        await facts.flush()
                        await event(facts, run)
                        await facts.commit()
                automation.data = {**d, "next_at": (now() + timedelta(seconds=max(60, int(d.get("interval_seconds", 3600))))).isoformat()}
            await db.commit()
        from .notifications import notify
        async with self.stores.facts() as facts:
            plans = (await facts.scalars(select(Record).where(Record.kind == "training-plans"))).all()
        for plan in plans:
            if plan.data.get("deleted") or plan.data.get("status") in {"completed", "cancelled"}:
                continue
            if plan.data.get("due_at", "9999") <= (now() + timedelta(days=2)).isoformat():
                await notify(self.stores, key=f"reminder:{plan.id}:{now().date().isoformat()}", org_id=plan.org_id, title="学习计划临近截止或已滞后", message="请进入学习中心查看课程和考核进度。", talent_id=plan.data.get("talent_id"), target="training-plans")


async def main():
    import signal
    import redis.asyncio as redis
    from redis.exceptions import RedisError

    worker = Worker()
    loop, task = asyncio.get_running_loop(), asyncio.current_task()
    for sig in (signal.SIGTERM, signal.SIGINT):
        with suppress(NotImplementedError):
            loop.add_signal_handler(sig, task.cancel)
    await worker.stores.initialize()
    worker.stores.redis = redis.from_url(worker.config.redis_url, decode_responses=True, socket_connect_timeout=2, socket_timeout=3)
    try:
        last = 0
        while True:
            if asyncio.get_running_loop().time() - last > 30:
                lock = worker.stores.redis.lock("talent:scheduler", timeout=60)
                acquired = False
                try:
                    acquired = await lock.acquire(blocking=False)
                    if acquired:
                        await worker.schedule()
                except RedisError:
                    # MySQL row locks plus deterministic PG run IDs remain the fallback.
                    await worker.schedule()
                finally:
                    if acquired:
                        with suppress(RedisError):
                            await lock.release()
                last = asyncio.get_running_loop().time()
            if not await worker.run_once():
                try:
                    await worker.stores.redis.blpop("talent:queue", timeout=1)
                except RedisError:
                    await asyncio.sleep(1)
    finally:
        await worker.stores.redis.aclose()
        await worker.stores.close()


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)
    asyncio.run(main())
