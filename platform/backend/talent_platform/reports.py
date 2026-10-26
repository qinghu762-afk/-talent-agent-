"""Safe reports built only from bounded query rows and locally drawn figures."""
from __future__ import annotations

import io
import json
import math
import re
from datetime import date
from decimal import Decimal, InvalidOperation
from html import escape
from pathlib import Path


def safe_text(value) -> str:
    text = value if isinstance(value, str) else json.dumps(value, ensure_ascii=False, default=str)
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff]", "", text)[:200000]


def chart_data(columns, rows):
    """Use actual rows, never model-supplied chart markup, URLs or data."""
    if len(columns) != 2 or not rows:
        return None
    data = []
    for row in rows:
        if len(row) != 2 or isinstance(row[1], bool) or row[1] is None:
            return None
        try:
            number = float(Decimal(str(row[1])))
        except (InvalidOperation, ValueError, TypeError, OverflowError):
            return None
        if not math.isfinite(number):
            return None
        data.append([safe_text(row[0]), number])
    if not math.isfinite(max(0, max(item[1] for item in data)) - min(0, min(item[1] for item in data))):
        return None
    temporal = []
    try:
        for label, _ in data:
            temporal.append(date.fromisoformat(label + '-01' if re.fullmatch(r'\d{4}-\d{2}', label) else label))
    except ValueError:
        temporal = []
    kind = 'line' if len(data) > 1 and temporal and temporal == sorted(set(temporal)) else 'bar'
    return {'type': kind, 'x': safe_text(columns[0]), 'y': safe_text(columns[1]), 'data': data[:30]}


def _chart_png(chart):
    from PIL import Image, ImageDraw, ImageFont
    font_paths = ('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc',
                  '/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc', 'C:/Windows/Fonts/msyh.ttc',
                  'C:/Windows/Fonts/simhei.ttf', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')
    font_path = next((path for path in font_paths if Path(path).is_file()), None)
    font = ImageFont.truetype(font_path, 26) if font_path else ImageFont.load_default(size=26)
    small = ImageFont.truetype(font_path, 22) if font_path else ImageFont.load_default(size=22)
    data = chart['data'][:12] if chart['type'] == 'bar' else chart['data']
    width, height = 1500, max(640, len(data) * 54 + 170) if chart['type'] == 'bar' else 720
    image = Image.new('RGB', (width, height), 'white')
    draw = ImageDraw.Draw(image)
    def fit(text, available, face):
        text = text[:500]
        if draw.textlength(text, font=face) <= available:
            return text
        while text and draw.textlength(text + '…', font=face) > available:
            text = text[:-1]
        return text + '…'
    draw.text((32, 24), fit(chart['y'] + ' / ' + chart['x'], width - 64, font), font=font, fill='#111827')
    values = [row[1] for row in data]
    low, high = min(0, min(values)), max(0, max(values))
    if low == high:
        high = low + 1
    if chart['type'] == 'bar':
        left, right, top, bottom = 390, 1280, 100, height - 70
        def x(value): return left + (value - low) / (high - low) * (right - left)
        for tick in range(6):
            value = low + (high - low) * tick / 5
            position = x(value)
            draw.line((position, top, position, bottom), fill='#e5e7eb', width=2)
            draw.text((position - 15, bottom + 15), f'{value:.3g}', font=small, fill='#4b5563')
        for index, (label, value) in enumerate(data):
            y = top + (index + .5) * (bottom - top) / len(data)
            label = fit(label, 340, font)
            draw.text((25, y - 16), label, font=font, fill='#111827')
            draw.rectangle((min(x(0), x(value)), y - 15, max(x(0), x(value)), y + 15), fill='#1d4ed8')
            draw.text((max(x(0), x(value)) + 10, y - 14), f'{value:.4g}', font=small, fill='#111827')
    else:
        left, right, top, bottom = 125, 1420, 110, 595
        points = []
        for tick in range(6):
            value = low + (high - low) * tick / 5
            y = bottom - (value - low) / (high - low) * (bottom - top)
            draw.line((left, y, right, y), fill='#e5e7eb', width=2)
            draw.text((15, y - 12), f'{value:.3g}', font=small, fill='#4b5563')
        for index, (label, value) in enumerate(data):
            x = left + index / max(1, len(data) - 1) * (right - left)
            y = bottom - (value - low) / (high - low) * (bottom - top)
            points.append((x, y))
            if index % max(1, math.ceil(len(data) / 6)) == 0 or index == len(data) - 1:
                draw.text((x, bottom + 35), fit(label, 140, small), font=small, fill='#111827', anchor='mm')
        draw.line(points, fill='#1d4ed8', width=5)
        for x, y in points:
            draw.ellipse((x - 5, y - 5, x + 5, y + 5), fill='#1d4ed8')
    output = io.BytesIO()
    image.save(output, format='PNG', dpi=(220, 220))
    return output.getvalue(), len(data)


def _table_rows(result):
    columns = list(result.get('columns') or [])
    if len(columns) > 20:
        raise ValueError('Report supports at most 20 columns')
    rows = [list(row) for row in (result.get('rows') or [])[:200]]
    table = [[safe_text(value) for value in columns]] if columns else []
    total = 0
    for row in rows:
        cells = [safe_text(row[index]) if index < len(row) else '' for index in range(len(columns))]
        total += sum(map(len, cells))
        if total > 200000:
            raise ValueError('Report table exceeds 200000 characters')
        table.append(cells)
    return columns, rows, table


def render_report(result: dict, format: str) -> bytes:
    if format not in {'docx', 'pdf'}:
        raise ValueError('Report format must be docx or pdf')
    columns, rows, table_data = _table_rows(result)
    chart = chart_data(columns, rows)
    figure, figure_count = _chart_png(chart) if chart else (None, 0)
    caption = f'图示使用查询结果前 {figure_count} 行；完整导出表最多 200 行。' if figure else ''
    title, output = '人才平台分析报告', io.BytesIO()
    sections = [('分析结论', result.get('answer', '')), ('查询语句', result.get('sql', '')),
                ('证据与说明', result.get('citations', result.get('feedback', [])))]
    sections = [item for item in sections if item[0] != '证据与说明' or item[1]]
    if format == 'docx':
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
        document = Document()
        section = document.sections[0]
        section.top_margin = section.bottom_margin = Inches(.75)
        section.left_margin = section.right_margin = Inches(.75)
        if len(columns) > 6:
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width, section.page_height = section.page_height, section.page_width
        for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
            style = document.styles[name]
            style.font.name = 'Microsoft YaHei' if Path('C:/Windows/Fonts/msyh.ttc').exists() else 'Noto Sans CJK SC'
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), style.font.name)
        document.styles['Normal'].font.size = Pt(10)
        document.add_paragraph(title, 'Title')
        document.add_heading(sections[0][0], level=1)
        for line in safe_text(sections[0][1]).splitlines() or ['']:
            document.add_paragraph(line)
        if figure:
            document.add_heading('结果图示', level=1)
            document.add_picture(io.BytesIO(figure), width=Inches(6.4))
            document.add_paragraph(caption)
        document.add_heading('查询结果', level=1)
        if table_data:
            table = document.add_table(rows=1, cols=len(columns))
            table.style = 'Table Grid'
            for index, label in enumerate(table_data[0]):
                table.rows[0].cells[index].text = label
            repeat = OxmlElement('w:tblHeader')
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row in table_data[1:]:
                for cell, value in zip(table.add_row().cells, row):
                    cell.text = value
            for row_index, row in enumerate(table.rows):
                for cell in row.cells:
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    properties = cell._tc.get_or_add_tcPr()
                    shade = OxmlElement('w:shd')
                    shade.set(qn('w:fill'), 'E2E8F0' if row_index == 0 else 'F8FAFC' if row_index % 2 == 0 else 'FFFFFF')
                    properties.append(shade)
                    borders = OxmlElement('w:tcBorders')
                    for edge in ('top', 'left', 'bottom', 'right'):
                        border = OxmlElement('w:' + edge)
                        for key, value in (('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')):
                            border.set(qn('w:' + key), value)
                        borders.append(border)
                    properties.append(borders)
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(4)
                        paragraph.paragraph_format.space_before = Pt(4)
                        if row_index == 0:
                            for run in paragraph.runs:
                                run.bold = True
        else:
            document.add_paragraph('本次查询没有返回明细。')
        for heading, value in sections[1:]:
            document.add_heading(heading, level=1)
            for line in safe_text(value).splitlines() or ['']:
                document.add_paragraph(line)
        document.save(output)
    else:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.platypus import Image, LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle
        pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
        styles = getSampleStyleSheet()
        for name in ('Title', 'Heading2', 'BodyText'):
            styles[name].fontName = 'STSong-Light'
            styles[name].textColor = colors.black
            styles[name].wordWrap = 'CJK'
        styles['BodyText'].leading = 15
        code_style = ParagraphStyle('Query', parent=styles['BodyText'], fontName='Courier', fontSize=8, leading=12)
        cell_style = ParagraphStyle('Cell', parent=styles['BodyText'], fontSize=9, leading=13, wordWrap='CJK', alignment=TA_LEFT)
        story = [Paragraph(title, styles['Title']), Spacer(1, 12), Paragraph('分析结论', styles['Heading2'])]
        def paragraphs(value):
            return [Paragraph(escape(line) or '&#160;', styles['BodyText']) for line in safe_text(value).splitlines() or ['']]
        story.extend(paragraphs(sections[0][1]))
        if figure:
            story.extend([Spacer(1, 12), Paragraph('结果图示', styles['Heading2'])])
            from PIL import Image as PILImage
            with PILImage.open(io.BytesIO(figure)) as image:
                image_height = 460 * image.height / image.width
            story.extend([Image(io.BytesIO(figure), width=460, height=image_height), Paragraph(caption, styles['BodyText'])])
        for heading, value in sections[1:]:
            body = [Paragraph(escape(line) or '&#160;', code_style) for line in safe_text(value).splitlines() or ['']] if heading == '查询语句' and safe_text(value).isascii() else paragraphs(value)
            story.extend([Spacer(1, 12), Paragraph(heading, styles['Heading2']), *body])
        story.extend([Spacer(1, 12), Paragraph('查询结果', styles['Heading2'])])
        page_size = landscape(A4) if len(columns) > 6 else A4
        if table_data:
            column_width = (page_size[0] - 90) / len(columns)
            # Bound physical row height before ReportLab pagination. Its splitInRow
            # can create repeated empty header rows after splitting a very tall cell.
            characters = max(12, int((column_width - 14) / 9) * 18)
            physical_rows = [table_data[0]]
            for row in table_data[1:]:
                fragments = [[value[start:start + characters] for start in range(0, len(value), characters)] or [''] for value in row]
                for index in range(max(map(len, fragments))):
                    values = [pieces[index] if index < len(pieces) else '' for pieces in fragments]
                    if index:
                        values[0] = '（续）' + values[0]
                    physical_rows.append(values)
            cells = [[Paragraph(escape(value).replace('\n', '<br/>') or '&#160;', cell_style) for value in row] for row in physical_rows]
            table = LongTable(cells, colWidths=[column_width] * len(columns), repeatRows=1, splitInRow=0)
            table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e2e8f0')),
                                      ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
                                      ('GRID', (0, 0), (-1, -1), .4, colors.HexColor('#d9d9d9')),
                                      ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                                      ('LEFTPADDING', (0, 0), (-1, -1), 7), ('RIGHTPADDING', (0, 0), (-1, -1), 7),
                                      ('TOPPADDING', (0, 0), (-1, -1), 6), ('BOTTOMPADDING', (0, 0), (-1, -1), 6)]))
            story.append(table)
        else:
            story.append(Paragraph('本次查询没有返回明细。', styles['BodyText']))
        def footer(canvas, document):
            canvas.setFont('STSong-Light', 9)
            canvas.drawRightString(page_size[0] - 45, 24, str(document.page))
        SimpleDocTemplate(output, title=title, author='人才平台', pagesize=page_size,
                          leftMargin=45, rightMargin=45, topMargin=40, bottomMargin=40).build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()
