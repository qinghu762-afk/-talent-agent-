"""Generate labelled synthetic parser fixtures without calling services or models.

Run with the existing project environment:
    .venv/Scripts/python.exe platform/tools/generate_acceptance_fixtures.py

The output must be new or empty. Existing material is never overwritten or removed.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
from pathlib import Path

NOTICE_ZH = '合成测试数据 仅供人才平台解析验收'
NOTICE_EN = 'SYNTHETIC TEST DATA FOR PARSER ACCEPTANCE'
FONT_PATHS = ('C:/Windows/Fonts/msyh.ttc', 'C:/Windows/Fonts/simsun.ttc',
              '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc',
              '/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc')


def _profile(number, language, scenario):
    english = language == 'en'
    person = {
        'name': f'Synthetic Candidate {number:02d}' if english else f'测试候选{number:02d}',
        'phone': f'+1 202 555 {100 + number:04d}',
        'email': f'candidate{number:02d}@example.test',
        'location': 'Shanghai' if english else '上海',
        'skills': ['Python', 'SQL'],
        'education': 'Bachelors degree in Computer Science' if english else '本科 计算机科学 合成示例大学',
        'work': 'Python Engineer at Synthetic Company 2020-01 to 2025-01' if english else '合成示例公司 Python工程师 2020-01 至 2025-01',
        'project': 'Procurement Automation System using Python and SQL' if english else '采购自动审批系统 使用Python和SQL开发',
    }
    missing = []
    if scenario == 'missing_contact':
        person['phone'] = person['email'] = None
        missing.append('contact')
    elif scenario == 'missing_name':
        person['name'] = None
        missing.append('name')
    elif scenario == 'missing_skills':
        person['skills'] = []
        person['work'] = 'Synthetic Company Engineer 2020-01 to 2025-01' if english else '合成示例公司 工程师 2020-01 至 2025-01'
        person['project'] = None
        missing.append('skills')
    elif scenario == 'missing_education_work':
        person['education'] = person['work'] = person['project'] = None
        missing.append('education_or_experiences')
    elif scenario == 'unknown_dates':
        person['work'] = 'Python Engineer at Synthetic Company; dates unavailable' if english else '合成示例公司 Python工程师 起止日期待确认'
    elif scenario == 'overlapping_dates':
        person['work'] += '\n' + ('Part-time SQL Engineer 2022-01 to 2024-01' if english else '兼职SQL工程师 2022-01 至 2024-01')
    elif scenario == 'email_only':
        person['phone'] = None
    return person, missing


def _rows(person, language, scenario):
    english = language == 'en'
    names = ('Name', 'Phone reserved for testing', 'Email', 'Location', 'Skills', 'Education', 'Experience', 'Project') if english else (
        '姓名', '约定测试电话', '邮箱', '现居地', '技能', '教育', '工作经历', '项目经历')
    values = (person['name'], person['phone'], person['email'], person['location'],
              ', '.join(person['skills']), person['education'], person['work'], person['project'])
    rows = [(label, value) for label, value in zip(names, values) if value]
    if scenario == 'negative_claims':
        rows.append(('边界说明', '没有PMP证书；不熟悉Java；Python和SQL用于实际项目。'))
    if scenario == 'reference_contact':
        rows.append(('推荐人 非候选人本人', '合成推荐人 联系电话 +1 202 555 0199；不得作为候选人联系方式。'))
    if language == 'mixed':
        rows.append(('English project note', 'Synthetic procurement automation project with Python and SQL.'))
    return rows


def _font(size):
    from PIL import ImageFont
    path = next((path for path in FONT_PATHS if Path(path).is_file()), None)
    if path is None:
        raise RuntimeError('A local CJK font is required; no font is downloaded automatically')
    return ImageFont.truetype(path, size)


def _wrap(text, fits):
    lines, line = [], ''
    for character in str(text):
        if character == '\n':
            lines.append(line)
            line = ''
        elif line and not fits(line + character):
            lines.append(line)
            line = character
        else:
            line += character
    if line:
        lines.append(line)
    return lines


def _scan_image(rows, language, *, rotate=False, low_contrast=False):
    from PIL import Image, ImageDraw
    image = Image.new('RGB', (1700, 2300), '#faf8f3' if low_contrast else 'white')
    draw = ImageDraw.Draw(image)
    heading, font = _font(39), _font(36)
    notice = NOTICE_EN if language == 'en' else NOTICE_ZH
    draw.text((90, 70), notice, font=heading, fill='#333333')
    y = 170
    for label, value in rows:
        for line in _wrap(f'{label}: {value}', lambda text: draw.textlength(text, font=font) <= 1510):
            draw.text((90, y), line, font=font, fill='#555555' if low_contrast else '#111111')
            y += 55
        y += 22
    if y > image.height - 80:
        raise ValueError('Synthetic scan content would be clipped')
    return image.rotate(2.5, expand=True, fillcolor='white') if rotate else image


def _docx(path, rows, language, layout):
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.oxml.ns import qn
    from docx.shared import Pt
    document = Document()
    normal = document.styles['Normal']
    normal.font.name = 'Microsoft YaHei'
    normal.font.size = Pt(11)
    normal.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    if layout == 'two_column':
        columns = document.sections[0]._sectPr.find(qn('w:cols'))
        columns.set(qn('w:num'), '2')
    document.add_heading(NOTICE_EN if language == 'en' else NOTICE_ZH, 0)
    if layout == 'table':
        table = document.add_table(rows=0, cols=2)
        table.style = 'Table Grid'
        for label, value in rows:
            cells = table.add_row().cells
            cells[0].text, cells[1].text = label, value
    else:
        for index, (label, value) in enumerate(rows):
            if layout == 'two_column' and index == 4:
                document.add_paragraph().add_run().add_break(WD_BREAK.COLUMN)
            if layout == 'multipage' and index in {4, 7}:
                document.add_page_break()
            document.add_heading(label, 1)
            document.add_paragraph(value)
    document.core_properties.title = 'Synthetic parser acceptance fixture'
    document.core_properties.subject = 'Synthetic data only'
    document.save(path)


def _native_pdf(path, rows, language, layout):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas
    pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
    pdf = canvas.Canvas(str(path), pagesize=A4)
    pdf.setTitle('Synthetic parser acceptance fixture')
    notice = NOTICE_EN if language == 'en' else NOTICE_ZH
    groups = [rows[:4], rows[4:7], rows[7:]] if layout == 'multipage' else [rows]
    groups = [group for group in groups if group]
    for group in groups:
        pdf.setFont('STSong-Light', 13)
        pdf.drawString(35, 800, notice)
        parts = [group[:4], group[4:]] if layout == 'two_column' else [group]
        for column, part in enumerate(parts):
            x, y = 35 + column * 280, 752
            available = 240 if layout == 'two_column' else 520
            pdf.setFont('STSong-Light', 11)
            for label, value in part:
                for line in _wrap(f'{label}: {value}', lambda text: pdfmetrics.stringWidth(text, 'STSong-Light', 11) <= available):
                    pdf.drawString(x, y, line)
                    y -= 21
                y -= 14
            if y < 45:
                raise ValueError('Synthetic native PDF content would be clipped')
        pdf.showPage()
    pdf.save()


def _scan_pdf(path, rows, language, layout):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    pdf = canvas.Canvas(str(path), pagesize=A4)
    pdf.setTitle('Synthetic image-only parser acceptance fixture')
    groups = [rows[:4], rows[4:]] if layout == 'multipage' else [rows]
    for group in groups:
        image = _scan_image(group, language)
        buffer = io.BytesIO()
        image.save(buffer, format='PNG')
        pdf.drawImage(ImageReader(buffer), 0, 0, width=A4[0], height=A4[1])
        pdf.showPage()
    pdf.save()


def _xlsx(path, rows, language, layout):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Synthetic profile'
    worksheet.append([NOTICE_ZH + ' / ' + NOTICE_EN])
    worksheet.append(['字段 Field', '原文 Source text'])
    groups = [rows[:4], rows[4:]] if layout == 'multiple_sheets' else [rows]
    for index, group in enumerate(groups):
        sheet = worksheet if index == 0 else workbook.create_sheet('Synthetic history')
        if index:
            sheet.append([NOTICE_ZH + ' / ' + NOTICE_EN])
            sheet.append(['字段 Field', '原文 Source text'])
        for row in group:
            sheet.append(list(row))
        sheet.column_dimensions['A'].width, sheet.column_dimensions['B'].width = 26, 95
        sheet.freeze_panes = 'A3'
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                cell.font = Font(name='Microsoft YaHei', size=11, bold=cell.row <= 2)
                if cell.row <= 2:
                    cell.fill = PatternFill('solid', fgColor='E2E8F0')
            sheet.row_dimensions[row[0].row].height = 35
    workbook.properties.title = 'Synthetic parser acceptance fixture'
    workbook.save(path)


def generate(output: Path | str) -> dict:
    output = Path(output)
    if output.is_symlink() or output.exists() and any(output.iterdir()):
        raise FileExistsError('Output must be a new or empty directory; existing material is preserved')
    _font(20)  # Check required local CJK glyph support before creating files.
    output.mkdir(parents=True, exist_ok=True)
    specs = [
        ('txt', 'zh', 'plain', 'complete'), ('txt', 'en', 'plain', 'complete'),
        ('txt', 'mixed', 'plain', 'complete'), ('txt', 'zh', 'plain', 'missing_contact'),
        ('txt', 'en', 'plain', 'missing_name'), ('txt', 'zh', 'plain', 'missing_skills'),
        ('txt', 'zh', 'plain', 'missing_education_work'), ('txt', 'zh', 'plain', 'unknown_dates'),
        ('txt', 'zh', 'plain', 'negative_claims'), ('txt', 'en', 'plain', 'overlapping_dates'),
        ('txt', 'zh', 'plain', 'reference_contact'), ('txt', 'zh', 'plain', 'noise'),
        ('txt', 'zh', 'utf8_bom', 'complete'), ('txt', 'en', 'plain', 'email_only'),
        ('docx', 'zh', 'table', 'complete'), ('docx', 'zh', 'two_column', 'complete'),
        ('docx', 'zh', 'multipage', 'complete'), ('docx', 'zh', 'table', 'missing_education_work'),
        ('docx', 'en', 'table', 'complete'), ('docx', 'mixed', 'two_column', 'complete'),
        ('pdf_native', 'zh', 'plain', 'complete'), ('pdf_native', 'zh', 'two_column', 'complete'),
        ('pdf_native', 'zh', 'multipage', 'complete'), ('pdf_native', 'en', 'two_column', 'complete'),
        ('pdf_native', 'zh', 'plain', 'missing_contact'), ('pdf_native', 'zh', 'plain', 'unknown_dates'),
        ('png_scan', 'zh', 'scan', 'complete'), ('jpg_scan', 'en', 'scan', 'complete'),
        ('png_scan', 'zh', 'skewed_scan', 'complete'), ('jpg_scan', 'zh', 'low_contrast_scan', 'complete'),
        ('pdf_scan', 'zh', 'scan', 'complete'), ('pdf_scan', 'en', 'multipage', 'complete'),
        ('pdf_scan', 'zh', 'scan', 'missing_education_work'),
        ('xlsx', 'zh', 'table', 'complete'), ('xlsx', 'en', 'multiple_sheets', 'complete'),
        ('pdf_corrupt', 'zh', 'broken', 'corrupt'),
    ]
    fixtures = []
    for number, (kind, language, layout, scenario) in enumerate(specs, 1):
        extension = {'pdf_native': 'pdf', 'pdf_scan': 'pdf', 'pdf_corrupt': 'pdf', 'png_scan': 'png', 'jpg_scan': 'jpg'}.get(kind, kind)
        filename = f'{number:02d}_synthetic_{language}_{layout}_{scenario}.{extension}'
        if not re.fullmatch(r'[a-zA-Z0-9_.-]+', filename) or Path(filename).name != filename:
            raise ValueError('Fixture names must be safe flat basenames')
        path = output / filename
        person, missing = _profile(number, language, scenario)
        rows = _rows(person, language, scenario)
        notice = NOTICE_EN if language == 'en' else NOTICE_ZH
        reference = notice + '\n' + '\n'.join(f'{label}: {value}' for label, value in rows)
        quality = 'partial' if missing else 'complete'
        if scenario == 'noise':
            reference = NOTICE_ZH + '\n字符噪声 @@@ ### ???\n本页不含任何人物档案。'
            rows, quality = [], 'fail'
        if kind == 'txt':
            path.write_text(reference + '\n', encoding='utf-8-sig' if layout == 'utf8_bom' else 'utf-8')
        elif kind == 'docx':
            _docx(path, rows, language, layout)
        elif kind == 'pdf_native':
            _native_pdf(path, rows, language, layout)
        elif kind in {'png_scan', 'jpg_scan'}:
            image = _scan_image(rows, language, rotate=layout == 'skewed_scan', low_contrast=layout == 'low_contrast_scan')
            image.save(path, format='PNG' if kind == 'png_scan' else 'JPEG', dpi=(200, 200), **({'quality': 86} if kind == 'jpg_scan' else {}))
        elif kind == 'pdf_scan':
            _scan_pdf(path, rows, language, layout)
        elif kind == 'xlsx':
            _xlsx(path, rows, language, layout)
        else:
            path.write_bytes(b'%PDF-1.7\n% SYNTHETIC CORRUPT FIXTURE\n1 0 obj << /Type /Catalog')
            quality, rows, reference = 'fail', [], 'Synthetic intentionally truncated PDF; no valid resume content'
        phrases = [value for label, value in rows if label in {'姓名', 'Name', '邮箱', 'Email', '技能', 'Skills', '教育', 'Education', '项目经历', 'Project'}]
        fixtures.append({'id': f'synthetic-{number:02d}', 'filename': filename, 'type': kind, 'language': language,
                         'layout': layout, 'scenario': scenario, 'synthetic': True, 'expected_quality': quality,
                         'expected_phrases': phrases, 'expected_missing_fields': missing,
                         'expected_profile': person if quality != 'fail' else {}, 'reference_text': reference,
                         'needs_review_expected': quality != 'complete' or scenario == 'unknown_dates',
                         'sha256': hashlib.sha256(path.read_bytes()).hexdigest()})
    manifest = {'version': 'synthetic-parser-fixtures-v1', 'synthetic': True, 'fixture_count': len(fixtures),
                'notice': '所有人物、机构、经历与联系方式均为合成验收数据，不代表真实候选人或实际业务。',
                'contact_policy': 'example.test mailboxes; NANP reserved fictional 202-555-0100 through 202-555-0199 telephone range',
                'expected_quality_semantics': 'Core-field presence label for source text, not a claim that OCR or model processing already passed. complete may still need review for unknown dates.',
                'limitations': ['No OCR, MinerU, model, API or Docker execution is performed.',
                                'DOCX files are parser inputs; native Office rendering is not part of generation.',
                                'Scans are digitally generated raster text, not real physical scans.',
                                'Corrupt PDF intentionally has no xref/trailer and must fail parsing.'],
                'fixtures': fixtures}
    (output / 'fixture-manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    return manifest


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, default=Path(__file__).resolve().parents[1] / 'verification' / 'fixtures')
    arguments = parser.parse_args()
    result = generate(arguments.output)
    print(json.dumps({'output': str(arguments.output.resolve()), 'fixtures': result['fixture_count'], 'synthetic': True}, ensure_ascii=False))
